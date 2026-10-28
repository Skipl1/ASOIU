"""Отчёты о пищевой ценности за период и их экспорт (PDF / XLSX / CSV)."""

import csv
import io
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import schemas
from app.config import get_settings
from app.models import RECIPE_ACCEPTED, FitnessGoal, Recipe, User
from app.services import nutrition
from app.services.profile import energy_for, get_profile

MAX_PERIOD_DAYS = 366

DEVIATION_STATUS = {"ok": "в норме", "below": "ниже цели", "above": "выше цели"}


def tz() -> ZoneInfo:
    return ZoneInfo(get_settings().app_timezone)


def local_dt(dt: datetime) -> datetime:
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(tz())


def local_date(dt: datetime) -> date:
    return local_dt(dt).date()


def validate_period(date_from: date, date_to: date) -> None:
    if date_to < date_from:
        raise HTTPException(422, detail="Дата окончания периода раньше даты начала")
    if (date_to - date_from).days + 1 > MAX_PERIOD_DAYS:
        raise HTTPException(422, detail=f"Период не может превышать {MAX_PERIOD_DAYS} дней")


def recipes_in_period(
    db: Session,
    user_id: uuid.UUID,
    date_from: date,
    date_to: date,
    goal_id: uuid.UUID | None = None,
    only_diet: bool = True,
) -> list[Recipe]:
    """Рецепты за период. Для отчётов и анализа (only_diet) — только включённые пользователем в рацион."""
    start = datetime.combine(date_from, time.min, tzinfo=tz()).astimezone(timezone.utc)
    end = datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=tz()).astimezone(timezone.utc)
    stmt = select(Recipe).where(Recipe.user_id == user_id, Recipe.created_at >= start, Recipe.created_at < end)
    if goal_id:
        stmt = stmt.where(Recipe.fitness_goal_id == goal_id)
    if only_diet:
        stmt = stmt.where(Recipe.generation_status == RECIPE_ACCEPTED)
    return list(db.scalars(stmt.order_by(Recipe.created_at)))


@dataclass
class ReportData:
    out: schemas.NutritionReportOut
    recipes: list[Recipe]
    user: User
    goal: FitnessGoal | None


def build_report(db: Session, user: User, date_from: date, date_to: date, goal_id: uuid.UUID | None = None) -> ReportData:
    validate_period(date_from, date_to)
    goal = db.get(FitnessGoal, goal_id) if goal_id else None
    if goal_id and goal is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Фитнес-ориентир не найден")
    recipes = recipes_in_period(db, user.user_id, date_from, date_to, goal_id)

    days: dict[date, dict] = {}
    d = date_from
    while d <= date_to:
        days[d] = {"calories": 0.0, "protein": 0.0, "fat": 0.0, "carbs": 0.0, "recipes": 0}
        d += timedelta(days=1)
    for r in recipes:
        day = days.get(local_date(r.created_at))
        if day is None:
            continue
        day["calories"] += r.calories_kcal
        day["protein"] += r.protein_total_g
        day["fat"] += r.fat_total_g
        day["carbs"] += r.carb_total_g
        day["recipes"] += 1

    total_cal = sum(v["calories"] for v in days.values())
    total_p = sum(v["protein"] for v in days.values())
    total_f = sum(v["fat"] for v in days.values())
    total_c = sum(v["carbs"] for v in days.values())
    days_with_data = sum(1 for v in days.values() if v["recipes"])
    avg = total_cal / days_with_data if days_with_data else 0.0
    pct = nutrition.macro_percent(total_p, total_f, total_c)

    target = deviation = None
    status_code = None
    profile = get_profile(db, user.user_id)
    if profile is not None:
        target = energy_for(profile, goal or profile.goal).target
        if days_with_data:
            deviation = avg - target
            rel = deviation / target
            status_code = "ok" if abs(rel) <= 0.1 else "below" if rel < 0 else "above"

    out = schemas.NutritionReportOut(
        date_from=date_from,
        date_to=date_to,
        goal_name=goal.name if goal else None,
        recipes_count=len(recipes),
        days_with_data=days_with_data,
        average_calories=round(avg, 1),
        total_calories=round(total_cal, 1),
        total_protein=round(total_p, 1),
        total_fat=round(total_f, 1),
        total_carbs=round(total_c, 1),
        macro_percent=schemas.NutritionTotals(calories=round(total_cal, 1), **pct),
        target_calories=round(target, 1) if target else None,
        deviation_calories=round(deviation, 1) if deviation is not None else None,
        deviation_status=DEVIATION_STATUS.get(status_code) if status_code else None,
        daily_calories=[
            schemas.DailyNutrition(
                date=k, calories=round(v["calories"], 1), protein=round(v["protein"], 1),
                fat=round(v["fat"], 1), carbs=round(v["carbs"], 1), recipes=v["recipes"],
            )
            for k, v in days.items()
        ],
    )
    return ReportData(out=out, recipes=recipes, user=user, goal=goal)


# ---------------- экспорт ----------------

def _fmt_date(d: date) -> str:
    return d.strftime("%d.%m.%Y")


def _summary_rows(r: schemas.NutritionReportOut) -> list[tuple[str, str]]:
    rows = [
        ("Период", f"{_fmt_date(r.date_from)} — {_fmt_date(r.date_to)}"),
        ("Фитнес-цель", r.goal_name or "все цели"),
        ("Рецептов в рационе за период", str(r.recipes_count)),
        ("Дней с данными", str(r.days_with_data)),
        ("Средняя калорийность дня, ккал", f"{r.average_calories:.0f}"),
        ("Калорийность за период, ккал", f"{r.total_calories:.0f}"),
        ("Белки за период, г", f"{r.total_protein:.1f}"),
        ("Жиры за период, г", f"{r.total_fat:.1f}"),
        ("Углеводы за период, г", f"{r.total_carbs:.1f}"),
        ("Доли Б/Ж/У, %", f"{r.macro_percent.protein:.1f} / {r.macro_percent.fat:.1f} / {r.macro_percent.carbs:.1f}"),
    ]
    if r.target_calories:
        rows.append(("Целевая калорийность дня, ккал", f"{r.target_calories:.0f}"))
    if r.deviation_calories is not None:
        rows.append(("Отклонение среднего дня, ккал", f"{r.deviation_calories:+.0f} ({r.deviation_status})"))
    return rows


def export_csv(data: ReportData) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Отчёт о пищевой ценности — АС ГРСН"])
    w.writerow(["Пользователь", data.user.name, data.user.email])
    for k, v in _summary_rows(data.out):
        w.writerow([k, v])
    w.writerow([])
    w.writerow(["Дата", "Рецептов", "Ккал", "Белки, г", "Жиры, г", "Углеводы, г"])
    for d in data.out.daily_calories:
        w.writerow([_fmt_date(d.date), d.recipes, f"{d.calories:.0f}", f"{d.protein:.1f}", f"{d.fat:.1f}", f"{d.carbs:.1f}"])
    w.writerow([])
    w.writerow(["Дата и время", "Рецепт", "Цель", "Ккал", "Белки, г", "Жиры, г", "Углеводы, г"])
    for r in data.recipes:
        w.writerow([
            local_dt(r.created_at).strftime("%d.%m.%Y %H:%M"), r.name, r.goal.name if r.goal else "",
            f"{r.calories_kcal:.0f}", f"{r.protein_total_g:.1f}", f"{r.fat_total_g:.1f}", f"{r.carb_total_g:.1f}",
        ])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(data: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="5B8CFF")

    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Отчёт о пищевой ценности — АС ГРСН"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{data.user.name} ({data.user.email})"
    for n, (k, v) in enumerate(_summary_rows(data.out), start=4):
        ws.cell(row=n, column=1, value=k).font = Font(bold=True)
        ws.cell(row=n, column=2, value=v)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 40

    def table(sheet, headers, rows, widths):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(row)
        for col, width in zip("ABCDEFG", widths):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A2"

    days = wb.create_sheet("По дням")
    table(
        days,
        ["Дата", "Рецептов", "Ккал", "Белки, г", "Жиры, г", "Углеводы, г"],
        [[d.date, d.recipes, d.calories, d.protein, d.fat, d.carbs] for d in data.out.daily_calories],
        [14, 10, 10, 10, 10, 12],
    )
    for cell in days["A"][1:]:
        cell.number_format = "DD.MM.YYYY"
    if data.out.daily_calories:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Калорийность по дням"
        chart.y_axis.title = "ккал"
        chart.x_axis.title = "Дата"
        n = len(data.out.daily_calories) + 1
        chart.add_data(Reference(days, min_col=3, min_row=1, max_row=n), titles_from_data=True)
        chart.set_categories(Reference(days, min_col=1, min_row=2, max_row=n))
        # openpyxl 3.1 не записывает delete=0 у осей, и Excel скрывает их: график выходит без подписей и делений
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.axPos, chart.y_axis.axPos = "b", "l"
        chart.x_axis.number_format = "DD.MM"
        chart.x_axis.majorTickMark = chart.y_axis.majorTickMark = "out"
        chart.x_axis.tickLblPos = chart.y_axis.tickLblPos = "low"
        chart.y_axis.scaling.min = 0
        chart.legend = None  # один ряд — легенда только отнимает место
        chart.gapWidth = 50
        chart.width, chart.height = 24, 10
        days.add_chart(chart, "H2")

    rec = wb.create_sheet("Рецепты")
    table(
        rec,
        ["Дата и время", "Рецепт", "Цель", "Ккал", "Белки, г", "Жиры, г", "Углеводы, г"],
        [[local_dt(r.created_at).replace(tzinfo=None), r.name, r.goal.name if r.goal else "",
          r.calories_kcal, r.protein_total_g, r.fat_total_g, r.carb_total_g] for r in data.recipes],
        [18, 48, 24, 10, 10, 10, 12],
    )
    for cell in rec["A"][1:]:
        cell.number_format = "DD.MM.YYYY HH:MM"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]


def _register_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "GRSN" in pdfmetrics.getRegisteredFontNames():
        return "GRSN", "GRSN-Bold"
    custom = get_settings().pdf_font_path
    candidates = ([(custom, custom)] if custom else []) + FONT_CANDIDATES
    for regular, bold in candidates:
        if regular and Path(regular).exists():
            pdfmetrics.registerFont(TTFont("GRSN", regular))
            pdfmetrics.registerFont(TTFont("GRSN-Bold", bold if Path(bold).exists() else regular))
            return "GRSN", "GRSN-Bold"
    raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Не найден шрифт с кириллицей для PDF (PDF_FONT_PATH)")


def export_pdf(data: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold = _register_fonts()
    title = ParagraphStyle("t", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h", fontName=bold, fontSize=12, leading=16, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("b", fontName=regular, fontSize=9.5, leading=12)

    def tbl(rows, widths, header=True):
        t = Table([[Paragraph(str(c), body) for c in row] for row in rows], colWidths=widths, repeatRows=1 if header else 0)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8B8B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6FF"))]
        t.setStyle(TableStyle(style))
        return t

    r = data.out
    story = [
        Paragraph("Отчёт о пищевой ценности рациона", title),
        Paragraph(f"АС ГРСН · {data.user.name} ({data.user.email}) · сформирован {datetime.now(tz()).strftime('%d.%m.%Y %H:%M')}", body),
        Spacer(1, 6 * mm),
        Paragraph("Сводные показатели", h2),
        tbl(_summary_rows(r), [80 * mm, 90 * mm], header=False),
        Paragraph("Показатели по дням", h2),
        tbl(
            [["Дата", "Рецептов", "Ккал", "Белки, г", "Жиры, г", "Углеводы, г"]]
            + [[_fmt_date(d.date), d.recipes, f"{d.calories:.0f}", f"{d.protein:.1f}", f"{d.fat:.1f}", f"{d.carbs:.1f}"]
               for d in r.daily_calories],
            [28 * mm, 22 * mm, 28 * mm, 28 * mm, 28 * mm, 30 * mm],
        ),
    ]
    if data.recipes:
        story += [
            Paragraph("Рецепты рациона за период", h2),
            tbl(
                [["Дата", "Рецепт", "Ккал", "Б / Ж / У, г"]]
                + [[local_dt(r_.created_at).strftime("%d.%m %H:%M"), r_.name, f"{r_.calories_kcal:.0f}",
                    f"{r_.protein_total_g:.0f} / {r_.fat_total_g:.0f} / {r_.carb_total_g:.0f}"] for r_ in data.recipes],
                [24 * mm, 96 * mm, 18 * mm, 32 * mm],
            ),
        ]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                      title="Отчёт о пищевой ценности", author="АС ГРСН").build(story)
    return buf.getvalue()


EXPORTERS = {"csv": export_csv, "xlsx": export_xlsx, "pdf": export_pdf}
MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
